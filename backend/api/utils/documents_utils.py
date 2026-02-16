# documents_router.py
from fastapi import HTTPException
from typing import List, Optional
from pathlib import Path
import io, json
import re
# --- Extractors ---
import pdfplumber
from docx import Document as DocxDocument
import openpyxl

UPLOAD_ROOT = Path("uploads")  # vervang later door object storage adapter
MAX_BYTES = 25 * 1024 * 1024   # 25MB cap als voorbeeld

ALLOWED_EXTS = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

SAFE_FILENAME_RE = re.compile(r"[^A-Za-z0-9._()-]+")


def sanitize_filename(name: str) -> str:
    name = name.strip().replace("\\", "_").replace("/", "_")
    name = SAFE_FILENAME_RE.sub("_", name)
    return name or "file"


def detect_by_extension(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in ALLOWED_EXTS:
        return ALLOWED_EXTS[ext]
    raise HTTPException(status_code=400, detail=f"Unsupported file extension: {ext}")


def decode_text_guess(data: bytes) -> str:
    # probeer utf-8, val terug op latin-1
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="ignore")


def extract_text(content: bytes, mime_type: str) -> str:
    bio = io.BytesIO(content)

    if mime_type == "text/plain":
        return decode_text_guess(content)

    if mime_type == "application/pdf":
        text_parts = []
        with pdfplumber.open(bio) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t:
                    text_parts.append(t)

                # Check of het een scan is (images-only)
                try:
                    if getattr(page, "images", None) and len(page.images) > 0:
                        has_any_images = True
                except Exception:
                    pass

        text = "\n".join(text_parts).strip()

        if not text:
            raise HTTPException(
                status_code=422,
                detail={
                    "type": "https://example.com/problems/scanned-pdf-unsupported",
                    "title": "Scanned PDF niet ondersteund",
                    "detail": (
                        "Dit PDF-bestand lijkt gescand (image-only) en bevat geen doorzoekbare tekst. "
                        "OCR wordt momenteel niet ondersteund. Upload een digitale PDF of voer eerst OCR uit."
                    ),
                    "status": 422,
                    "code": "SCANNED_PDF_UNSUPPORTED",
                    "has_images": has_any_images,
                },
            )

        return text

    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = DocxDocument(bio)
        paras = [p.text for p in doc.paragraphs if p.text]
        return "\n".join(paras).strip()

    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        wb = openpyxl.load_workbook(bio, read_only=True, data_only=True)
        out_lines = []
        for sheet in wb.worksheets:
            out_lines.append(f"### Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                # maak nette tab-gescheiden regels
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    out_lines.append("\t".join(cells))
        return "\n".join(out_lines).strip()

    raise HTTPException(status_code=400, detail="Unsupported file type")


def parse_tags(tags_raw: Optional[str]) -> List[str]:
    if not tags_raw:
        return []
    # probeer JSON-lijst; val terug op CSV
    try:
        parsed = json.loads(tags_raw)
        if isinstance(parsed, list):
            return [str(x).strip() for x in parsed if str(x).strip()]
    except json.JSONDecodeError:
        pass
    return [t.strip() for t in tags_raw.split(",") if t.strip()]